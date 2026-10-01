#!/usr/bin/env python3
"""Registry-driven, read-only RAW -> staging runner for any configured shop.

Safety boundary:
- never opens or writes the production Data Mart;
- never modifies UI artifacts;
- every selected shop gets an independent QA report and artifact namespace;
- exit code is non-zero when any selected shop fails QA.
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import os
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import Any, Dict, Iterable, List, Tuple

from google.oauth2.service_account import Credentials
from googleapiclient.discovery import build
from googleapiclient.http import MediaIoBaseDownload
from openpyxl import load_workbook

from modules.ads_product_mart import aggregate_product_ads_rows
from modules.catalog_resolver import ListingRecord
from modules.multi_shop_staging import (
    audit_listing_catalog_snapshot,
    audit_bi_workbook_schema,
    build_listing_records,
    normalize_bi_shop_daily,
    normalize_bi_traffic,
    normalize_listing_sales_info_export,
    normalize_orders,
    normalize_product_performance,
    parse_ads_csv,
    qa_staging,
    run_catalog_resolver,
)
from modules.schema_registry import SchemaContractError, require_schema, summarize_schema_audits
from modules.shop_registry import (
    CORE_STAGING_DOMAINS,
    enabled_shops,
    load_shop_registry,
    select_shops,
)

DRIVE_SCOPE="https://www.googleapis.com/auth/drive.readonly"


def drive_client():
    raw=os.environ.get("GOOGLE_SERVICE_ACCOUNT_JSON","")
    if not raw:
        raise RuntimeError("GOOGLE_SERVICE_ACCOUNT_JSON is required")
    creds=Credentials.from_service_account_info(json.loads(raw),scopes=[DRIVE_SCOPE])
    return build("drive","v3",credentials=creds,cache_discovery=False)


def children(drive,folder_id):
    out=[]; token=None
    while True:
        r=drive.files().list(
            q=f"'{folder_id}' in parents and trashed=false",
            fields="nextPageToken,files(id,name,mimeType,modifiedTime,size)",
            pageSize=1000,pageToken=token,orderBy="name",
            supportsAllDrives=True,includeItemsFromAllDrives=True).execute()
        out.extend(r.get("files",[])); token=r.get("nextPageToken")
        if not token:
            return out


def child_folder(drive,parent_id,name,required=True):
    xs=[x for x in children(drive,parent_id)
        if x.get("name")==name and x.get("mimeType")=="application/vnd.google-apps.folder"]
    if len(xs)==1:
        return xs[0]["id"]
    if required:
        raise FileNotFoundError(f"Expected one folder {name!r}; found {len(xs)}")
    return ""


def _candidate_root_id(drive,candidate):
    typ=str(candidate.get("type") or "")
    if typ=="folder_id":
        return str(candidate.get("folder_id") or "")
    if typ=="child_folder":
        return child_folder(
            drive,
            str(candidate.get("parent_folder_id") or ""),
            str(candidate.get("folder_name") or ""),
            required=False,
        )
    return ""


def resolve_shop_raw_root(drive,shop):
    """Select the highest-priority configured root containing all core domains."""
    diagnostics=[]
    candidates=sorted(shop.get("raw_root_candidates") or [],key=lambda x:int(x.get("priority",0)),reverse=True)
    required=set(CORE_STAGING_DOMAINS)
    for candidate in candidates:
        root_id=_candidate_root_id(drive,candidate)
        diag={
            "type":candidate.get("type"),
            "priority":int(candidate.get("priority",0)),
            "folder_id":root_id,
            "folder_name":candidate.get("folder_name",""),
        }
        if not root_id:
            diag.update({"status":"NOT_FOUND","missing_domains":sorted(required)})
            diagnostics.append(diag)
            continue
        names={x.get("name") for x in children(drive,root_id)
               if x.get("mimeType")=="application/vnd.google-apps.folder"}
        missing=sorted(required-names)
        diag.update({"status":"READY" if not missing else "INCOMPLETE","missing_domains":missing})
        diagnostics.append(diag)
        if not missing:
            return root_id,{
                "status":"READY",
                "selected":diag,
                "candidates":diagnostics,
            }
    raise RuntimeError(
        f"{shop.get('shop_key')}: no configured raw root contains all core domains; "
        f"candidates={diagnostics}"
    )


def download_bytes(drive,file_id):
    req=drive.files().get_media(fileId=file_id)
    buf=io.BytesIO(); dl=MediaIoBaseDownload(buf,req,chunksize=1024*1024); done=False
    while not done:
        _,done=dl.next_chunk()
    return buf.getvalue()


def _first_worksheet_path(zf):
    """Resolve the first worksheet through workbook relationships.

    Shopee exports can use sheet2.xml as the only visible worksheet, so
    hard-coding sheet1.xml is not safe.
    """
    main_ns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_attr="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    pkg_rel_ns="http://schemas.openxmlformats.org/package/2006/relationships"

    workbook=ET.fromstring(zf.read("xl/workbook.xml"))
    sheet=workbook.find(f"{{{main_ns}}}sheets/{{{main_ns}}}sheet")
    if sheet is None:
        raise ValueError("XLSX workbook has no worksheet")

    rid=sheet.attrib.get(f"{{{rel_attr}}}id","")
    if not rid:
        raise ValueError("XLSX first worksheet relationship is missing")

    rels=ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    target=""
    for rel in rels.findall(f"{{{pkg_rel_ns}}}Relationship"):
        if rel.attrib.get("Id")==rid:
            target=rel.attrib.get("Target","")
            break
    if not target:
        raise ValueError("XLSX first worksheet target is missing")

    target=target.lstrip("/")
    if target.startswith("xl/"):
        path=posixpath.normpath(target)
    else:
        path=posixpath.normpath(posixpath.join("xl",target))
    if path not in zf.namelist():
        raise ValueError(f"XLSX first worksheet not found: {path}")
    return path


def xlsx_first_sheet_rows_tolerant(blob):
    """Read first worksheet directly from XML, ignoring broken dimension metadata."""
    with zipfile.ZipFile(io.BytesIO(blob)) as zf:
        ns={"x":"http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
        shared=[]
        if "xl/sharedStrings.xml" in zf.namelist():
            root=ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in root.findall("x:si",ns):
                shared.append("".join(
                    (t.text or "")
                    for t in si.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t")
                ))

        sheet_path=_first_worksheet_path(zf)
        root=ET.fromstring(zf.read(sheet_path))
        out=[]
        for row in root.findall(".//x:sheetData/x:row",ns):
            cells={}; max_col=-1
            for cell in row.findall("x:c",ns):
                ref=cell.attrib.get("r","A1")
                m=re.match(r"([A-Z]+)",ref)
                if not m:
                    continue
                col=0
                for ch in m.group(1):
                    col=col*26+(ord(ch)-64)
                col-=1; max_col=max(max_col,col)
                typ=cell.attrib.get("t","")
                if typ=="inlineStr":
                    node=cell.find("x:is",ns)
                    val="".join(
                        (t.text or "")
                        for t in node.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}t")
                    ) if node is not None else ""
                else:
                    node=cell.find("x:v",ns)
                    raw="" if node is None or node.text is None else node.text
                    val=shared[int(raw)] if typ=="s" and raw else raw
                cells[col]=val
            vals=[""]*(max_col+1)
            for col,val in cells.items():
                vals[col]=val
            out.append(vals)
        return out


def xlsx_sheets(blob):
    wb=load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    try:
        return [[list(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets]
    finally:
        wb.close()


def domain_month(drive,root_id,domain,month):
    did=child_folder(drive,root_id,domain)
    yid=child_folder(drive,did,month[:4])
    return child_folder(drive,yid,month)


def one_file(drive,folder_id,suffix):
    xs=[x for x in children(drive,folder_id) if x.get("name","").lower().endswith(suffix)]
    if len(xs)!=1:
        raise FileNotFoundError(f"Expected one *{suffix} file; found {len(xs)}")
    return xs[0]


def load_listing_catalog(drive,root_id,month,shop):
    if not bool(shop.get("listing_catalog_enabled",True)):
        return [],[],{"enabled":False,"authority":"product_performance_fallback"}
    try:
        mid=domain_month(drive,root_id,"listing_catalog",month)
    except FileNotFoundError:
        return [],[],{"enabled":True,"folderPresent":False,"files":[],"authority":"product_performance_fallback"}
    fid=child_folder(drive,mid,"final",required=False) or mid
    files=[x for x in children(drive,fid) if x.get("name","").lower().endswith(".xlsx")]
    if not files:
        return [],[],{"enabled":True,"folderPresent":True,"files":[],"authority":"product_performance_fallback"}
    latest=sorted(files,key=lambda x:(x.get("modifiedTime",""),x.get("name","")))[-1]
    rows=xlsx_first_sheet_rows_tolerant(download_bytes(drive,latest["id"]))
    if not rows:
        raise ValueError("Listing catalog workbook has no rows")
    snapshot_at=latest.get("modifiedTime") or ""
    products,variations,meta=normalize_listing_sales_info_export(
        rows,
        expected_shopee_shop_id=str(shop["shopee_shop_id"]),
        shop_id=str(shop["shop_id"]),
        snapshot_at=snapshot_at,
        source_name=latest["name"],
    )
    audit=audit_listing_catalog_snapshot(products,variations)
    if audit["status"]!="PASS":
        raise ValueError(f"Listing catalog snapshot QA failed: {audit}")
    return products,variations,{
        "enabled":True,
        "folderPresent":True,
        "authority":"seller_center_mass_update_sales_info",
        "file":{"id":latest["id"],"name":latest["name"],"modifiedTime":latest.get("modifiedTime")},
        "meta":meta,
        "audit":audit,
    }


def load_product_performance(drive,root_id,month,shop_id):
    mid=domain_month(drive,root_id,"product_performance",month)
    f=one_file(drive,child_folder(drive,mid,"final"),".xlsx")
    sheets=xlsx_sheets(download_bytes(drive,f["id"]))
    if not sheets:
        raise ValueError("Product Performance workbook has no sheets")
    schema_audit=require_schema(
        "product_performance",
        sheets[0][0] if sheets[0] else [],
        source_name=f["name"],
    )
    p,v=normalize_product_performance(sheets[0],shop_id=shop_id,data_month=month)
    return p,v,{"id":f["id"],"name":f["name"],"modifiedTime":f.get("modifiedTime"),
                "schema_audit":schema_audit}


def load_bi(drive,root_id,month,shop_id):
    mid=domain_month(drive,root_id,"business_insights",month)
    f=one_file(drive,child_folder(drive,mid,"final"),".xlsx")
    sheets=xlsx_sheets(download_bytes(drive,f["id"]))
    schema_audits=audit_bi_workbook_schema(sheets,source_name=f["name"])
    shop_rows=normalize_bi_shop_daily(sheets,shop_id=shop_id)
    td,tm=normalize_bi_traffic(sheets,shop_id=shop_id,data_month=month)
    return shop_rows,td,tm,{"id":f["id"],"name":f["name"],"sheetCount":len(sheets),
        "modifiedTime":f.get("modifiedTime"),"schema_audits":schema_audits}


def select_orders_source_folder(drive,mid,month):
    """Prefer latest immutable snapshot, then legacy final/, then month root."""
    items=children(drive,mid)
    snaps=[x for x in items
           if x.get("mimeType")=="application/vnd.google-apps.folder"
           and x.get("name","").startswith("snapshot_")]
    if snaps:
        return sorted(snaps,key=lambda x:x["name"])[-1]
    finals=[x for x in items
            if x.get("mimeType")=="application/vnd.google-apps.folder"
            and x.get("name")=="final"]
    if len(finals)>1:
        raise ValueError(f"{month}: multiple Orders final folders")
    return finals[0] if finals else {"id":mid,"name":month}


def load_orders(drive,root_id,month,shop_id,loaded_at):
    mid=domain_month(drive,root_id,"orders",month)
    src=select_orders_source_folder(drive,mid,month)
    files=[x for x in children(drive,src["id"]) if x.get("name","").lower().endswith(".xlsx")]
    if not files:
        raise FileNotFoundError("No Orders XLSX found")
    orders={}; items=[]; manifest=[]
    for f in sorted(files,key=lambda x:x["name"]):
        blob=download_bytes(drive,f["id"])
        sheets=xlsx_sheets(blob)
        rows=sheets[0] if sheets else []
        header=[str(x or "").strip() for x in (rows[0] if rows else [])]
        if "Mã đơn hàng" not in header or "Ngày đặt hàng" not in header:
            rows=xlsx_first_sheet_rows_tolerant(blob)
        schema_audit=require_schema(
            "orders",
            rows[0] if rows else [],
            source_name=f["name"],
        )
        os_,it=normalize_orders(rows,shop_id=shop_id,loaded_at=loaded_at)
        for o in os_:
            oid=o["order_id"]
            if oid in orders and orders[oid]!=o:
                raise ValueError(f"Order {oid} differs across parts")
            orders[oid]=o
        items.extend(it)
        manifest.append({"id":f["id"],"name":f["name"],"modifiedTime":f.get("modifiedTime"),
                         "schema_audit":schema_audit})
    seq={}
    for x in items:
        oid=x["order_id"]; seq[oid]=seq.get(oid,0)+1; x["order_item_seq"]=seq[oid]
    return list(orders.values()),items,{"folder":src["name"],"files":manifest}


def load_ads(drive,root_id,month,shop):
    try:
        mid=domain_month(drive,root_id,"ads",month)
    except FileNotFoundError:
        return [],[],{"folderPresent":False,"files":[]}
    files=[x for x in children(drive,mid) if x.get("name","").lower().endswith(".csv")]
    rows=[]; meta=[]
    for f in sorted(files,key=lambda x:x["name"]):
        text=download_bytes(drive,f["id"]).decode("utf-8-sig")
        part,m=parse_ads_csv(
            text,
            expected_shopee_shop_id=str(shop["shopee_shop_id"]),
            shop_id=str(shop["shop_id"]),
            source_name=f["name"],
        )
        rows.extend(part)
        meta.append({**m,"file_id":f["id"],"file_name":f["name"],"modifiedTime":f.get("modifiedTime")})
    return rows,meta,{
        "folderPresent":True,
        "files":[{"id":x["id"],"name":x["name"],"modifiedTime":x.get("modifiedTime")} for x in files],
    }


def load_reference_file(path):
    if not path:
        return []
    rows=json.loads(Path(path).read_text(encoding="utf-8"))
    return [
        ListingRecord(
            shop_id=str(x.get("shop_id","")),
            product_id=str(x.get("product_id","")),
            title=str(x.get("title","")),
            parent_sku=str(x.get("parent_sku","")),
            variation_skus=set(x.get("variation_skus") or []),
        )
        for x in rows
    ]


def load_map(path):
    return json.loads(Path(path).read_text(encoding="utf-8")) if path else {}


def write_jsonl(path:Path,rows:Iterable[Dict[str,Any]]):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open("w",encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r,ensure_ascii=False,sort_keys=True,default=str)+"\n")


def load_catalog_evidence(drive,shop,month):
    root,root_resolution=resolve_shop_raw_root(drive,shop)
    listing_products,listing_variations,listing_src=load_listing_catalog(drive,root,month,shop)
    if listing_products:
        products,variations=listing_products,listing_variations
        authority="seller_center_listing_catalog"
        fallback_src=None
    else:
        products,variations,fallback_src=load_product_performance(
            drive,root,month,str(shop["shop_id"])
        )
        authority="product_performance_fallback"
    return {
        "shop":shop,
        "root_id":root,
        "root_resolution":root_resolution,
        "products":products,
        "variations":variations,
        "records":build_listing_records(products,variations),
        "authority":authority,
        "listing_source":listing_src,
        "fallback_source":fallback_src,
    }


def build_reference_pool(drive,registry,target_shop_key,month,cache):
    refs=[]; evidence=[]
    for other in enabled_shops(registry):
        if other["shop_key"]==target_shop_key:
            continue
        key=other["shop_key"]
        try:
            if key not in cache:
                cache[key]=load_catalog_evidence(drive,other,month)
            item=cache[key]
            refs.extend(item["records"])
            evidence.append({
                "shop_key":key,
                "shop_id":other["shop_id"],
                "status":"READY",
                "authority":item["authority"],
                "records":len(item["records"]),
            })
        except Exception as exc:
            evidence.append({
                "shop_key":key,
                "shop_id":other["shop_id"],
                "status":"SKIPPED",
                "error":str(exc),
            })
    return refs,evidence


def stage_shop(drive,registry,shop,month,output_dir,reference_file,family_map,catalog_cache):
    shop_key=str(shop["shop_key"]); sid=str(shop["shop_id"])
    now=dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat()
    root,root_resolution=resolve_shop_raw_root(drive,shop)

    products,variations,pp_src=load_product_performance(drive,root,month,sid)
    listing_products,listing_variations,listing_src=load_listing_catalog(drive,root,month,shop)
    bi,traffic_d,traffic_m,bi_src=load_bi(drive,root,month,sid)
    orders,items,orders_src=load_orders(drive,root,month,sid,now)
    ads,ads_meta,ads_src=load_ads(drive,root,month,shop)

    catalog_products=listing_products or products
    catalog_variations=listing_variations or variations
    catalog_cache[shop_key]={
        "shop":shop,
        "root_id":root,
        "root_resolution":root_resolution,
        "products":catalog_products,
        "variations":catalog_variations,
        "records":build_listing_records(catalog_products,catalog_variations),
        "authority":"seller_center_listing_catalog" if listing_products else "product_performance_fallback",
        "listing_source":listing_src,
        "fallback_source":None if listing_products else pp_src,
    }
    auto_refs,reference_evidence=build_reference_pool(
        drive,registry,shop_key,month,catalog_cache
    )
    references=auto_refs+reference_file
    catalog=run_catalog_resolver(
        build_listing_records(catalog_products,catalog_variations),
        references,
        family_map,
    )
    product_ads=aggregate_product_ads_rows(ads,expected_shop_id=sid)

    schema_audits=[]
    if listing_src.get("meta",{}).get("schema_audit"):
        schema_audits.append(listing_src["meta"]["schema_audit"])
    if pp_src.get("schema_audit"):
        schema_audits.append(pp_src["schema_audit"])
    schema_audits.extend(bi_src.get("schema_audits") or [])
    for file_meta in orders_src.get("files") or []:
        if file_meta.get("schema_audit"):
            schema_audits.append(file_meta["schema_audit"])
    for file_meta in ads_meta:
        sa=file_meta.get("schema_audit") or {}
        if sa.get("metadata"):
            schema_audits.append(sa["metadata"])
        if sa.get("rows"):
            schema_audits.append(sa["rows"])
    schema_summary=summarize_schema_audits(schema_audits)

    qa=qa_staging(
        shop_id=sid,
        target_month=month,
        orders=orders,
        order_items=items,
        products=products,
        variations=variations,
        bi_shop=bi,
        ads=ads,
        ads_files=ads_meta,
        catalog_resolutions=catalog,
        traffic_daily=traffic_d,
        traffic_monthly=traffic_m,
    )
    qa["checks"].append({
        "name":"source_schema_guard",
        "status":"PASS" if schema_summary["status"]!="FAIL" else "FAIL",
        "detail":schema_summary,
    })
    if schema_summary["status"]=="FAIL":
        qa["status"]="FAIL"
        qa["production_write_allowed"]=False
    qa.update({
        "schema_guard":schema_summary,
        "generated_at":now,
        "shop_key":shop_key,
        "display_name":shop.get("display_name"),
        "raw_root_resolution":root_resolution,
        "reference_catalogs":reference_evidence,
        "source_manifest":{
            "listing_catalog":listing_src,
            "product_performance":pp_src,
            "business_insights":bi_src,
            "orders":orders_src,
            "ads":ads_src,
        },
        "safety":{
            "output_mode":"LOCAL_STAGING_ARTIFACT_ONLY",
            "production_data_mart_opened":False,
            "production_data_mart_written":False,
            "ui_modified":False,
        },
    })

    out=Path(output_dir)/shop_key/month
    for name,rows in {
        "fact_orders":orders,
        "fact_order_items":items,
        "fact_product_performance_monthly":products,
        "fact_product_variations_monthly":variations,
        "fact_shop_performance_daily":bi,
        "fact_traffic_source_daily":traffic_d,
        "fact_traffic_source_monthly":traffic_m,
        "fact_ads_performance_daily":ads,
        "dm_ads_product_daily_candidate":product_ads,
        "catalog_resolution_staging":catalog,
        "listing_catalog_products_snapshot":listing_products,
        "listing_catalog_variations_snapshot":listing_variations,
    }.items():
        write_jsonl(out/f"{name}.jsonl",rows)
    out.mkdir(parents=True,exist_ok=True)
    (out/"staging_qa_report.json").write_text(
        json.dumps(qa,ensure_ascii=False,indent=2,sort_keys=True,default=str),
        encoding="utf-8",
    )
    (out/"schema_drift_report.json").write_text(
        json.dumps({
            "shop_key":shop_key,
            "shop_id":sid,
            "month":month,
            "summary":schema_summary,
            "audits":schema_audits,
        },ensure_ascii=False,indent=2,sort_keys=True,default=str),
        encoding="utf-8",
    )
    return qa


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument(
        "--shop-key",
        action="append",
        default=[],
        help="Registry shop key. Repeat to stage multiple shops. Omit to stage all enabled shops.",
    )
    ap.add_argument("--month",required=True)
    ap.add_argument("--registry",default="config/shop_registry.json")
    ap.add_argument("--output-dir",default="staging_artifacts")
    ap.add_argument("--reference-catalog-json",default="")
    ap.add_argument("--sku-family-map-json",default="")
    a=ap.parse_args()

    if not re.fullmatch(r"\d{4}-\d{2}",a.month):
        raise ValueError("--month must be YYYY-MM")

    registry=load_shop_registry(a.registry)
    selected=select_shops(registry,a.shop_key or ["all"])
    drive=drive_client()
    external_refs=load_reference_file(a.reference_catalog_json)
    family_map=load_map(a.sku_family_map_json)
    catalog_cache={}
    results=[]

    for shop in selected:
        try:
            qa=stage_shop(
                drive,registry,shop,a.month,a.output_dir,
                external_refs,family_map,catalog_cache,
            )
            schema_guard=qa.get("schema_guard") or {}
            results.append({
                "shop_key":shop["shop_key"],
                "shop_id":shop["shop_id"],
                "status":qa["status"],
                "production_write_allowed":qa["production_write_allowed"],
                "schema_status":schema_guard.get("status",""),
                "schema_warn_count":int((schema_guard.get("status_counts") or {}).get("WARN",0)),
                "schema_new_fingerprint_count":len(schema_guard.get("new_or_unbaselined_fingerprints") or []),
                "schema_unknown_column_groups":len(schema_guard.get("unknown_columns") or []),
                "error":"",
            })
        except Exception as exc:
            error_record={
                "shop_key":shop["shop_key"],
                "shop_id":shop["shop_id"],
                "month":a.month,
                "status":"ERROR",
                "production_write_allowed":False,
                "schema_status":"FAIL" if isinstance(exc,SchemaContractError) else "",
                "schema_audit":exc.audit if isinstance(exc,SchemaContractError) else {},
                "error":str(exc),
                "generated_at":dt.datetime.now(dt.timezone.utc).replace(microsecond=0).isoformat(),
                "safety":{
                    "production_data_mart_written":False,
                    "ui_modified":False,
                },
            }
            error_dir=Path(a.output_dir)/str(shop["shop_key"])/a.month
            error_dir.mkdir(parents=True,exist_ok=True)
            (error_dir/"staging_error.json").write_text(
                json.dumps(error_record,ensure_ascii=False,indent=2,sort_keys=True),
                encoding="utf-8",
            )
            results.append(error_record)

    overall="PASS" if results and all(x["status"]=="PASS" for x in results) else "FAIL"
    summary={
        "status":overall,
        "month":a.month,
        "selected_shop_count":len(selected),
        "results":results,
        "safety":{
            "production_data_mart_written":False,
            "ui_modified":False,
        },
    }
    root=Path(a.output_dir)
    root.mkdir(parents=True,exist_ok=True)
    (root/f"multi_shop_staging_summary_{a.month}.json").write_text(
        json.dumps(summary,ensure_ascii=False,indent=2,sort_keys=True),
        encoding="utf-8",
    )
    print(json.dumps(summary,ensure_ascii=False))
    return 0 if overall=="PASS" else 2


if __name__=="__main__":
    raise SystemExit(main())
